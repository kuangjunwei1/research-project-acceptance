import importlib.util
import io
import json
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from pypdf import PdfWriter

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('inspect_materials',ROOT/'scripts/inspect_materials.py')
mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)

class InspectTests(unittest.TestCase):
 def setUp(self):self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
 def pdf(self,name='0501-test.pdf',pages=2):
  p=self.root/name;w=PdfWriter()
  for _ in range(pages):w.add_blank_page(width=595,height=842)
  with p.open('wb')as f:w.write(f)
  return p
 def test_pdf_count(self):self.assertEqual(mod.metadata(self.pdf())['pdf_pages'],2)
 def test_valid_pages(self):self.assertEqual(mod.page_check(self.pdf(),[1,2])['out_of_range'],[])
 def test_out_of_range(self):self.assertEqual(mod.page_check(self.pdf(),[3])['out_of_range'],[3])
 def test_zero_rejected(self):
  with self.assertRaises(ValueError):mod.page_check(self.pdf(),[0])
 def test_negative_rejected(self):
  with self.assertRaises(ValueError):mod.page_check(self.pdf(),[-1])
 def test_empty_pages_rejected(self):
  with self.assertRaises(ValueError):mod.page_check(self.pdf(),[])
 def test_exact_duplicates(self):
  (self.root/'a.txt').write_text('same');(self.root/'b.txt').write_text('same')
  self.assertEqual(mod.inventory(self.root)['exact_duplicates'],[['a.txt','b.txt']])
 def test_different_content_not_duplicate(self):
  (self.root/'a.txt').write_text('one');(self.root/'b.txt').write_text('two')
  self.assertEqual(mod.inventory(self.root)['exact_duplicates'],[])
 def test_number_collision(self):
  (self.root/'1001-a.txt').write_text('a');(self.root/'1001-b.txt').write_text('b')
  self.assertIn('1001',mod.inventory(self.root)['number_collisions'])
 def test_word_pdf_pair_not_collision(self):
  self.pdf('0301-work.pdf');(self.root/'0301-work.docx').write_bytes(b'broken')
  self.assertEqual(mod.inventory(self.root)['number_collisions'],{})
 def test_invalid_pdf_recorded(self):
  (self.root/'broken.pdf').write_bytes(b'not a pdf')
  self.assertIn('error',mod.inventory(self.root)['files'][0])
 def test_invalid_docx_recorded(self):
  (self.root/'broken.docx').write_bytes(b'not zip')
  self.assertIn('error',mod.inventory(self.root)['files'][0])
 def test_invalid_xlsx_recorded(self):
  (self.root/'broken.xlsx').write_bytes(b'not zip')
  self.assertIn('error',mod.inventory(self.root)['files'][0])
 def test_docx_container_metadata(self):
  p=self.root/'test.docx'
  with zipfile.ZipFile(p,'w')as z:z.writestr('word/document.xml','<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p/><w:tbl><w:tr><w:tc><w:p/></w:tc></w:tr></w:tbl></w:body></w:document>')
  self.assertEqual(mod.metadata(p)['paragraphs'],2);self.assertEqual(mod.metadata(p)['tables'],1)
 def test_xlsx_container_metadata(self):
  p=self.root/'test.xlsx'
  with zipfile.ZipFile(p,'w')as z:z.writestr('xl/workbook.xml','<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheets><sheet name="成果"/></sheets></workbook>')
  self.assertEqual(mod.metadata(p)['sheet_names'],['成果'])
 def test_input_bytes_unchanged(self):
  p=self.pdf();before=p.read_bytes();mod.inventory(self.root);mod.page_check(p,[1]);self.assertEqual(p.read_bytes(),before)
 def test_unsupported_format_explicit(self):
  p=self.root/'old.xls';p.write_bytes(b'sample')
  self.assertEqual(mod.metadata(p)['metadata_status'],'not_inspected')
 def test_cli_error_exit(self):
  f=io.StringIO()
  with redirect_stdout(f):rc=mod.main(['page-check',str(self.pdf()),'--pages','3'])
  self.assertEqual(rc,2);self.assertEqual(json.loads(f.getvalue())['out_of_range'],[3])
 def test_cli_success_exit(self):
  f=io.StringIO()
  with redirect_stdout(f):rc=mod.main(['inventory',str(self.root)])
  self.assertEqual(rc,0);self.assertEqual(json.loads(f.getvalue())['files'],[])
 def test_missing_directory_exit(self):
  with redirect_stdout(io.StringIO()):rc=mod.main(['inventory',str(self.root/'missing')])
  self.assertEqual(rc,2)
 def test_malformed_page_exit(self):
  with redirect_stdout(io.StringIO()):rc=mod.main(['page-check',str(self.pdf()),'--pages','one'])
  self.assertEqual(rc,2)
 def test_real_docx(self):
  from docx import Document
  d=Document();d.add_paragraph('合成报告');d.add_table(rows=2,cols=2);p=self.root/'real.docx';d.save(p)
  before=p.read_bytes();self.assertEqual(mod.metadata(p)['tables'],1);self.assertEqual(p.read_bytes(),before)

if __name__=='__main__':unittest.main(verbosity=2)
